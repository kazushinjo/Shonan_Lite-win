"""Windows版操作説明書(DOCX)の生成スクリプト(日本語版・英語版)。

章データはapp/gui/manual_content_win.py(日本語)とmanual_content_win_en.py(英語)を
アプリ内Helpと共有する。実行すると次の2ファイルをこのスクリプトと同じフォルダへ出力する:
  - shonan_lite_win_operation_manual.docx     (日本語)
  - shonan_lite_win_operation_manual_en.docx  (English)

必要: python-docx (pip install python-docx)
"""
from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'gui'))
from manual_content_win import (  # noqa: E402
    MANUAL_DATE, MANUAL_SCREENSHOTS_WIN, MANUAL_SECTIONS_WIN, MANUAL_VERSION)
from manual_content_win_en import MANUAL_SCREENSHOTS_WIN_EN, MANUAL_SECTIONS_WIN_EN  # noqa: E402

IMAGE_DIR = ROOT / 'docs' / 'images'

LANGUAGES = {
    'ja': {
        'out': 'shonan_lite_win_operation_manual.docx',
        'sections': MANUAL_SECTIONS_WIN,
        'screenshots': MANUAL_SCREENSHOTS_WIN,
        'subtitle': 'DVB-S2 DATV送受信システム 操作説明書（Windows版）',
        'cover_info': ('対象: Windows 10/11 (x64) + ADALM-Pluto\n'
                       '接続先: Pluto（IPアドレスは出力設定画面で指定）\n'
                       f'版: v{MANUAL_VERSION} / {MANUAL_DATE}'),
        'toc': '目次',
        'note': ('本書はShonan_Lite for Windows v' + MANUAL_VERSION + 'のソースコードの挙動に基づいて記述しています。'
                 'ご使用の環境によって動作が異なる場合や、未発見の不具合が含まれる可能性があります。'),
        'figure': '図: ',
        'appendix': '付録A Pi 5実機版との主な相違点',
        'appendix_head': ('項目', 'Windows版の挙動'),
        'appendix_rows': (
            ('起動時のPluto再起動', '行わない(すぐホーム画面を表示)。'),
            ('Langstone V3・プリセット・Pluto電源', '非対応。ホーム画面に表示しない。'),
            ('ホーム画面右下2枠', '「アプリ再起動」「アプリ終了」(Pluto電源オフではない)。'),
            ('送信と受信の同時実行', 'オンデバイス復調がOFFの場合は排他制御(同時実行不可)。起動のたびにOFFへ戻る。'),
            ('受信復調', '別途radioconda(GNU Radio)が必要。公式インストーラが自動導入。'),
            ('映像ソース', 'ffmpeg dshowで列挙したWindowsのカメラ/キャプチャデバイス。'),
            ('システム日時設定', 'Linux専用のため動作しない。Windows自体の時計設定を使う。'),
            ('オンスクリーンキーボード', 'Windows版には組み込まれていない。物理キーボード+IMEで入力する。'),
            ('ウィンドウ', '通常のWindowsウィンドウ(大きさ変更可)。'),
        ),
        'credit': ('受信部の方式考案・受信部原システム設計: 山崎慎慈氏(JE1BTA) '
                   'rpi-dvbs2-receiver-guiの設計に基づきます\n'
                   '受信部安定化調査修正・再捕捉修正・本アプリ開発: 真城和一(JA6FUF/JH1XHX)\n'
                   '本アプリは、Dave Crump氏(G8GKQ)が開発したDATV送受信機プロジェクト「Portsdown」に啓発され、'
                   '開発したものです。同氏の先駆的な取り組みに感謝いたします。\n'
                   '本プログラムを使用して生じたいかなる損害についても、開発者は一切の責任を負いません。'
                   'ご自身の責任においてご利用ください。'),
        'header': 'Shonan_Lite for Windows 操作説明書',
        'east_asia_font': 'Yu Gothic',
    },
    'en': {
        'out': 'shonan_lite_win_operation_manual_en.docx',
        'sections': MANUAL_SECTIONS_WIN_EN,
        'screenshots': MANUAL_SCREENSHOTS_WIN_EN,
        'subtitle': 'DVB-S2 DATV Transceiver System Operation Manual (Windows Edition)',
        'cover_info': ('Target: Windows 10/11 (x64) + ADALM-Pluto\n'
                       'Destination: Pluto (set the IP address on the Stream Output screen)\n'
                       f'Version: v{MANUAL_VERSION} / {MANUAL_DATE}'),
        'toc': 'Contents',
        'note': ('This manual is written from the behavior of the source code of Shonan_Lite for Windows v'
                 + MANUAL_VERSION + '. Behavior may differ depending on your environment, and undiscovered '
                 'defects may remain.'),
        'figure': 'Figure: ',
        'appendix': 'Appendix A Main Differences from the Pi 5 Unit',
        'appendix_head': ('Item', 'Behavior of the Windows edition'),
        'appendix_rows': (
            ('Pluto restart at startup', 'Not performed (the Home screen is shown immediately).'),
            ('Langstone V3, Presets, Pluto Power', 'Not supported. Not shown on the Home screen.'),
            ('Bottom-right two cards on Home', '"App Restart" and "Exit App" (not a Pluto power-off).'),
            ('Simultaneous TX and RX', 'Exclusive when on-device demodulation is OFF (cannot run at the same time). It returns to OFF at every startup.'),
            ('Receive demodulation', 'Requires radioconda (GNU Radio) separately. The official installer installs it automatically.'),
            ('Video source', 'Windows camera/capture devices enumerated by ffmpeg dshow.'),
            ('System date/time setting', 'Linux-only, so it does not work. Use the Windows clock settings.'),
            ('On-screen keyboard', 'Not included in the Windows edition. Use a physical keyboard + IME.'),
            ('Window', 'A normal Windows window (resizable).'),
        ),
        'credit': ('Receiver method and original receiver system design: Shinji Yamazaki (JE1BTA), '
                   'based on rpi-dvbs2-receiver-gui.\n'
                   'Receiver stabilization, reacquisition fixes, and application development: '
                   'Kazuichi Shinjo (JA6FUF/JH1XHX).\n'
                   'This application was developed inspired by "Portsdown", the DATV transceiver project created '
                   'by Dave Crump (G8GKQ). We extend our deep gratitude for his pioneering work.\n'
                   'The developers accept no liability whatsoever for any damage arising from the use of this '
                   'program. Use it at your own risk.'),
        'header': 'Shonan_Lite for Windows Operation Manual',
        'east_asia_font': None,
    },
}


def set_east_asia_font(style, name):
    if not name:
        return
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    fonts.set(qn('w:eastAsia'), name)


def build(cfg):
    d = Document()
    sec = d.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    sec.different_first_page_header_footer = True

    styles = d.styles
    normal = styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    set_east_asia_font(normal, cfg['east_asia_font'])
    for name, size, color, before, after in (
            ('Heading 1', 16, '2E74B5', 18, 10), ('Heading 2', 13, '2E74B5', 14, 7),
            ('Heading 3', 12, '1F4D78', 10, 5)):
        s = styles[name]
        s.font.name = 'Calibri'
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        set_east_asia_font(s, cfg['east_asia_font'])

    # Cover page
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(120)
    r = p.add_run('Shonan_Lite for Windows')
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor(31, 77, 120)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(cfg['subtitle'])
    r.bold = True
    r.font.size = Pt(18)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.add_run(cfg['cover_info']).font.size = Pt(11)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    r = p.add_run(cfg['note'])
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(100, 100, 100)
    d.add_page_break()

    d.add_heading(cfg['toc'], level=1)
    for title, items in cfg['sections']:
        d.add_paragraph(title, style='List Number')
        for subtitle, _ in items:
            q = d.add_paragraph(subtitle, style='List Bullet 2')
            q.paragraph_format.space_after = Pt(2)
    d.add_page_break()

    for title, items in cfg['sections']:
        d.add_heading(title, level=1)
        image_names = cfg['screenshots'].get(title, [])
        if isinstance(image_names, str):
            image_names = [image_names]
        for image_name in image_names:
            image_path = IMAGE_DIR / image_name
            if not image_path.exists():
                continue
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(image_path), width=Inches(6.0))
            cap = d.add_paragraph(cfg['figure'] + title)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].italic = True
            cap.runs[0].font.size = Pt(9)
            cap.runs[0].font.color.rgb = RGBColor(100, 100, 100)
        for subtitle, text in items:
            d.add_heading(subtitle, level=2)
            d.add_paragraph(text)

    d.add_heading(cfg['appendix'], level=1)
    t = d.add_table(rows=1, cols=2)
    t.style = 'Table Grid'
    t.autofit = False
    t.columns[0].width = Inches(1.875)
    t.columns[1].width = Inches(4.625)
    for c, v in zip(t.rows[0].cells, cfg['appendix_head']):
        c.text = v
    for a, b in cfg['appendix_rows']:
        cells = t.add_row().cells
        cells[0].text = a
        cells[1].text = b

    first_footer = sec.first_page_footer.paragraphs[0]
    first_footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    first_footer.paragraph_format.line_spacing = 1.0
    credit = first_footer.add_run(cfg['credit'])
    credit.font.size = Pt(8)
    credit.font.color.rgb = RGBColor(100, 100, 100)

    header = sec.header.paragraphs[0]
    header.text = cfg['header']
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Page ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    out = Path(__file__).with_name(cfg['out'])
    d.save(out)
    print(out)


if __name__ == '__main__':
    for cfg in LANGUAGES.values():
        build(cfg)
